"""Render structured, verified resume data. stdin JSON -> stdout PDF/DOCX bytes."""
import io
import json
import os
import sys
from pathlib import Path
from xml.sax.saxutils import escape

data = json.load(sys.stdin)
resume = data['resume']

SECTION_LABELS = {
    'education': '教育经历',
    'skills': '专业技能',
    'projects': '项目经历',
    'experience': '工作与实习经历',
    'awards': '荣誉与证书',
}
DEFAULT_ORDER = ['education', 'skills', 'projects', 'experience', 'awards']
REQUIRED_TEMPLATES = {'campus-project', 'experienced', 'minimal'}


def text(value):
    return str(value or '').replace('\x00', '')


def list_items(key):
    values = [text(v).strip() for v in resume.get(key, []) if text(v).strip()]
    return [' · '.join(values)] if key == 'skills' and values else values


def load_templates():
    config_file = Path(__file__).with_name('resume_layouts.json')
    loaded = json.loads(config_file.read_text(encoding='utf-8'))
    templates = loaded.get('templates')
    if not isinstance(templates, dict) or not REQUIRED_TEMPLATES.issubset(templates):
        raise ValueError('resume_layouts.json 缺少必要模板。')
    for name, template in templates.items():
        if not isinstance(template, dict):
            raise ValueError(f'简历模板 {name} 配置无效。')
    return templates


def auto_template():
    counts = {key: len([v for v in resume.get(key, []) if text(v).strip()]) for key in DEFAULT_ORDER}
    total_chars = sum(len(text(v)) for key in DEFAULT_ORDER for v in resume.get(key, [])) + len(text(resume.get('summary')))
    if counts['experience'] >= max(2, counts['projects']) or sum(len(text(v)) for v in resume.get('experience', [])) >= 360:
        return 'experienced'
    if total_chars < 900 and counts['projects'] <= 2 and counts['experience'] <= 1:
        return 'minimal'
    return 'campus-project'


def normalize_layout():
    templates = load_templates()
    raw = data.get('layout') or resume.get('layout') or {}
    if isinstance(raw, str):
        raw = {'template': raw}
    if not isinstance(raw, dict):
        raw = {}
    template_name = raw.get('template') if raw.get('template') in templates else auto_template()
    template = {**templates[template_name]}
    requested_order = raw.get('sectionOrder') or template.get('sectionOrder') or DEFAULT_ORDER
    seen = []
    for key in requested_order:
        if key in SECTION_LABELS and key not in seen:
            seen.append(key)
    for key in DEFAULT_ORDER:
        if key not in seen:
            seen.append(key)
    template['name'] = template_name
    template['sectionOrder'] = seen
    return template


layout = normalize_layout()
contact = '  ·  '.join(text(v) for v in resume.get('contact', {}).values() if v)
heading = text(resume.get('targetTitle') or resume.get('title'))
output = io.BytesIO()


def safe_dimension(value, fallback):
    try:
        return float(value)
    except (TypeError, ValueError):
        return fallback


def safe_margin(name, fallback):
    margin = layout.get('margin') if isinstance(layout.get('margin'), dict) else {}
    return safe_dimension(margin.get(name), fallback)


if data['format'] == 'pdf':
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether
    from reportlab.lib import colors
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.lib.styles import ParagraphStyle
    requested_font = os.environ.get('RESUME_PROTOCOL_FONT', '')
    font_path = next((candidate for candidate in [
        requested_font,
        'C:/Windows/Fonts/msyh.ttc',
        '/System/Library/Fonts/STHeiti Medium.ttc',
        '/System/Library/Fonts/Hiragino Sans GB.ttc',
        '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
    ] if candidate and os.path.isfile(candidate)), '')
    if font_path:
        pdfmetrics.registerFont(TTFont('Resume', font_path, subfontIndex=0))
        font = 'Resume'
    else:
        pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
        font = 'STSong-Light'

    body_size = max(10.5, safe_dimension(layout.get('bodySize'), 10.5))
    leading = max(body_size + 4, safe_dimension(layout.get('leading'), 16))
    item_gap = safe_dimension(layout.get('itemGap'), 5)
    section_gap = safe_dimension(layout.get('sectionGap'), 8)
    keep_chars = int(safe_dimension(layout.get('paragraphKeepChars'), 180))
    body = ParagraphStyle(
        'Body',
        fontName=font,
        fontSize=body_size,
        leading=leading,
        textColor=colors.HexColor('#252525'),
        wordWrap='CJK',
        splitLongWords=1,
        spaceAfter=item_gap,
    )
    title = ParagraphStyle('Title', parent=body, fontSize=safe_dimension(layout.get('titleSize'), 21), leading=29, spaceAfter=5)
    sub = ParagraphStyle('Sub', parent=body, fontSize=10.5, leading=15, textColor=colors.HexColor('#555555'), spaceAfter=3)
    section = ParagraphStyle(
        'Section',
        parent=body,
        fontSize=safe_dimension(layout.get('sectionSize'), 12.5),
        leading=17,
        textColor=colors.HexColor('#111111'),
        spaceBefore=section_gap,
        spaceAfter=4,
        borderWidth=0,
        borderPadding=0,
    )

    def para(value, style=body):
        value = escape(text(value)).replace('\n', '<br/>')
        return Paragraph(value, style)

    def flow_item(value):
        return para(value)

    story = [para(resume.get('name') or '个人简历', title)]
    left_margin = safe_margin('left', 42)
    right_margin = safe_margin('right', 42)
    top_margin = safe_margin('top', 34)
    bottom_margin = safe_margin('bottom', 42)
    usable_width = 595.28 - left_margin - right_margin
    if heading:
        story.append(para(heading, sub))
    if resume.get('sourceNotice'):
        story.append(para(resume['sourceNotice'], sub))
    if contact:
        story.append(para(contact, sub))
    if resume.get('summary'):
        story.extend([Spacer(1, 6), flow_item(resume['summary'])])
    for key in layout['sectionOrder']:
        items = list_items(key)
        if not items:
            continue
        first = items[0]
        if len(text(first)) <= keep_chars:
            story.append(KeepTogether([para(SECTION_LABELS[key], section), para(first)]))
        else:
            first_para = para(first)
            first_parts = first_para.split(usable_width, leading * 2.2)
            if len(first_parts) > 1:
                story.append(KeepTogether([para(SECTION_LABELS[key], section), first_parts[0]]))
                story.extend(first_parts[1:])
            else:
                story.append(para(SECTION_LABELS[key], section))
                story.append(first_para)
        story.extend(flow_item(v) for v in items[1:])

    def footer(c, doc):
        c.saveState()
        c.setFont(font, 8)
        c.setFillColor(colors.HexColor('#777777'))
        c.drawRightString(550, 25, str(doc.page))
        c.restoreState()

    SimpleDocTemplate(
        output,
        pagesize=(595.28, 841.89),
        rightMargin=right_margin,
        leftMargin=left_margin,
        topMargin=top_margin,
        bottomMargin=bottom_margin,
        title=text(resume.get('name')) + ' - 简历',
        author='',
    ).build(story, onFirstPage=footer, onLaterPages=footer)
else:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def enable_word_wrap(paragraph):
        p_pr = paragraph._p.get_or_add_pPr()
        if p_pr.find(qn('w:wordWrap')) is None:
            p_pr.append(OxmlElement('w:wordWrap'))

    def apply_run_style(run, size=None, bold=False, color='252525'):
        run.font.name = 'Microsoft YaHei'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        run.font.size = Pt(size or max(10.5, safe_dimension(layout.get('bodySize'), 10.5)))
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)

    def paragraph(value, size=None, bold=False, color='252525', space_after=None, keep_together=True, keep_next=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after if space_after is not None else safe_dimension(layout.get('itemGap'), 5))
        p.paragraph_format.line_spacing = max(1.18, safe_dimension(layout.get('leading'), 16) / max(10.5, safe_dimension(layout.get('bodySize'), 10.5)))
        p.paragraph_format.keep_together = keep_together
        p.paragraph_format.keep_with_next = keep_next
        enable_word_wrap(p)
        for i, part in enumerate(text(value).split('\n')):
            if i:
                p.add_run().add_break()
            apply_run_style(p.add_run(part), size=size, bold=bold, color=color)
        return p

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.top_margin = Inches(safe_margin('top', 34) / 72)
    sec.bottom_margin = Inches(safe_margin('bottom', 42) / 72)
    sec.left_margin = Inches(safe_margin('left', 42) / 72)
    sec.right_margin = Inches(safe_margin('right', 42) / 72)
    style = doc.styles['Normal']
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(max(10.5, safe_dimension(layout.get('bodySize'), 10.5)))
    style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    style.paragraph_format.space_after = Pt(safe_dimension(layout.get('itemGap'), 5))
    style.paragraph_format.line_spacing = 1.25

    paragraph(resume.get('name') or '个人简历', size=safe_dimension(layout.get('titleSize'), 21), bold=True, color='111111', space_after=4, keep_together=True)
    if heading:
        paragraph(heading, size=10.5, color='555555', space_after=2)
    if resume.get('sourceNotice'):
        paragraph(resume['sourceNotice'], size=10.5, color='555555', space_after=2)
    if contact:
        paragraph(contact, size=10.5, color='555555', space_after=5)
    if resume.get('summary'):
        paragraph(resume['summary'], keep_together=len(text(resume['summary'])) <= int(safe_dimension(layout.get('paragraphKeepChars'), 180)))
    for key in layout['sectionOrder']:
        items = list_items(key)
        if not items:
            continue
        paragraph(SECTION_LABELS[key], size=safe_dimension(layout.get('sectionSize'), 12.5), bold=True, color='111111', space_after=3, keep_together=True, keep_next=True)
        for value in items:
            paragraph(value, keep_together=len(text(value)) <= int(safe_dimension(layout.get('paragraphKeepChars'), 180)))
    doc.save(output)

sys.stdout.buffer.write(output.getvalue())
